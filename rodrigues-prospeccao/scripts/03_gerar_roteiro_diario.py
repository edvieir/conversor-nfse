#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
GERADOR DO ROTEIRO DIARIO DE PROSPECCAO
Rodrigues Contabilidade

Le o CRM_Prospeccao_Rodrigues.xlsx e monta o roteiro do dia:
    1. Quem respondeu           -> prioridade maxima
    2. Follow-ups atrasados     -> zerar a fila
    3. Follow-ups de hoje       -> D3, D7 ou D30
    4. Novos primeiros contatos -> 10 D0 ja escritos e personalizados

Sai em roteiro_do_dia.txt, pronto para copiar e colar mensagem por mensagem.

Uso:
    python 03_gerar_roteiro_diario.py
    python 03_gerar_roteiro_diario.py --novos 15
    python 03_gerar_roteiro_diario.py --marcar     (grava o D0 de hoje no CRM)

Requisitos: pip install pandas openpyxl numpy
"""

import argparse
import datetime as dt
import os
import sys

import numpy as np
import pandas as pd
from openpyxl import load_workbook

CRM = "crm/CRM_Prospeccao_Rodrigues.xlsx"
SAIDA = "saidas/roteiro_do_dia.txt"
SEU_NOME = "Rodrigo"          # <-- TROQUE PELO SEU PRIMEIRO NOME

# Produtos monofasicos citados na mensagem, por segmento
PRODUTOS = {
    "autopec": "pneus, filtros e oleos lubrificantes",
    "pneu": "pneus e camaras de ar",
    "bebida": "cerveja, refrigerante e agua",
    "cosmetic": "perfumaria, cosmeticos e higiene pessoal",
    "farmac": "medicamentos e produtos de higiene pessoal",
    "minimercado": "bebidas, produtos de higiene e cosmeticos",
    "mercearia": "bebidas, produtos de higiene e cosmeticos",
    "combustiv": "combustiveis e lubrificantes",
    "lubrific": "oleos lubrificantes",
}
PRODUTO_PADRAO = "bebidas, cosmeticos e produtos de higiene pessoal"


def produtos_de(segmento):
    s = (segmento or "").lower()
    for chave, valor in PRODUTOS.items():
        if chave in s:
            return valor
    return PRODUTO_PADRAO


def saudacao():
    h = dt.datetime.now().hour
    return "Bom dia" if h < 12 else ("Boa tarde" if h < 18 else "Boa noite")


# =============================================================================
# MENSAGENS
# =============================================================================

def d0(linha):
    empresa = linha["Nome_Fantasia"] or linha["Razao_Social"]
    cunha = (linha["Cunha"] or "Monofasico").strip().lower()

    if "fator" in cunha:
        return (
            f"{saudacao()}. {SEU_NOME}, da Rodrigues Contabilidade.\n\n"
            f"Pergunta objetiva sobre a {empresa}: sua empresa esta no "
            f"Anexo III ou no Anexo V do Simples?\n\n"
            f"Se a folha, incluindo pro-labore, chega a 28% da receita dos "
            f"ultimos 12 meses, a aliquota inicial cai de 15,5% para 6%. "
            f"Vejo muita empresa parada a 2 ou 3 pontos desse limite sem "
            f"ninguem ajustar.\n\n"
            f"Com dois dados - faturamento e folha dos ultimos 12 meses - "
            f"eu calculo seu Fator R e te mando o numero. Sem custo e sem "
            f"compromisso.\n\n"
            f"Quer que eu faca?"
        )

    if "recem" in cunha:
        return (
            f"{saudacao()}. {SEU_NOME}, da Rodrigues Contabilidade.\n\n"
            f"Vi que a {empresa} foi constituida ha pouco tempo. Parabens.\n\n"
            f"Nos primeiros 12 meses tres coisas costumam sair errado e "
            f"custam caro depois: enquadramento tributario definido no "
            f"chute, pro-labore mal dimensionado e obrigacoes acessorias "
            f"entregues fora do prazo.\n\n"
            f"Preparei um checklist de 12 itens do que precisa estar "
            f"resolvido nesse primeiro ano. Te mando sem compromisso, mesmo "
            f"que voce ja tenha contador.\n\n"
            f"Posso enviar?"
        )

    prod = produtos_de(linha["Segmento"])
    return (
        f"{saudacao()}. Aqui e {SEU_NOME}, da Rodrigues Contabilidade.\n\n"
        f"Estou falando com estabelecimentos como a {empresa} por um motivo "
        f"bem especifico: a maior parte paga PIS e COFINS duas vezes sobre "
        f"os mesmos produtos.\n\n"
        f"Itens como {prod} ja tem esses tributos recolhidos na industria. "
        f"Quando nao sao separados no PGDAS, a empresa recolhe de novo - "
        f"todo mes, ha anos.\n\n"
        f"Consigo verificar isso no seu caso em cerca de 20 minutos, com os "
        f"PGDAS dos ultimos 3 meses. Se nao houver nada, eu te falo direto "
        f"e encerro o assunto.\n\n"
        f"Posso olhar?"
    )


def d3(linha):
    return (
        "Fechei uma revisao aqui em Fortaleza esses dias: R$ 1.840 por mes "
        "de PIS/COFINS pagos indevidamente, mais o retroativo dos ultimos "
        "5 anos.\n\n"
        "E exatamente o que eu queria checar no seu caso. Continua de pe."
    )


def d7(linha):
    return (
        "Nao quero insistir. Se nao for o momento, me responde so "
        "'agora nao' que eu paro por aqui.\n\n"
        "Se for, e so mandar os PGDAS que eu cuido do resto."
    )


def d30(linha):
    return (
        "Passando so para deixar um alerta util, sem oferta nenhuma:\n\n"
        "[ESCREVA AQUI: prazo que esta vencendo, mudanca de aliquota, "
        "novo limite do Simples, obrigacao nova]\n\n"
        "Se precisar de alguma coisa, e so chamar."
    )


# =============================================================================
# LEITURA E CLASSIFICACAO
# =============================================================================

def carregar():
    if not os.path.exists(CRM):
        sys.exit(f"[ERRO] '{CRM}' nao encontrado nesta pasta.")

    df = pd.read_excel(CRM, sheet_name="Pipeline", header=1, engine="openpyxl")
    df = df[df["Razao_Social"].notna()].copy()
    df = df[~df["Observacoes"].astype(str).str.contains("EXEMPLO", na=False)]

    for c in ["Razao_Social", "Nome_Fantasia", "Segmento", "Cunha",
              "Status", "Telefone", "Bairro", "Respondeu"]:
        df[c] = df[c].fillna("").astype(str).str.strip()

    df["Data_D0"] = pd.to_datetime(df["Data_D0"], errors="coerce")
    return df


def dia_util(data, dias):
    return pd.Timestamp(
        np.busday_offset(data.date(), dias, roll="forward")
    )


def classificar(df):
    hoje = pd.Timestamp(dt.date.today())
    encerrados = {"FECHADO", "Perdido", "Descartado"}

    responderam, atrasados, hoje_lista, novos = [], [], [], []

    for _, l in df.iterrows():
        if l["Status"] in encerrados:
            continue

        if l["Respondeu"].lower() == "sim":
            responderam.append(l)
            continue

        if pd.isna(l["Data_D0"]):
            novos.append(l)
            continue

        d0_data = l["Data_D0"]
        marcos = [("D3", dia_util(d0_data, 3), d3),
                  ("D7", dia_util(d0_data, 7), d7),
                  ("D30", dia_util(d0_data, 30), d30)]

        pendente = None
        for etapa, data, func in marcos:
            if data <= hoje and l["Status"] not in (f"{etapa} enviado",):
                pendente = (etapa, data, func)
        if pendente is None:
            continue

        etapa, data, func = pendente
        item = (l, etapa, func)
        (atrasados if data < hoje else hoje_lista).append(item)

    return responderam, atrasados, hoje_lista, novos


# =============================================================================
# SAIDA
# =============================================================================

def bloco(l, titulo, texto):
    zap = "".join(filter(str.isdigit, l["Telefone"]))
    link = f"https://wa.me/55{zap}" if len(zap) == 11 else "(fixo - ligar)"
    return (
        f"\n{'-' * 68}\n"
        f"{titulo}\n"
        f"{l['Razao_Social']}  |  {l['Bairro']}  |  {l['Telefone']}\n"
        f"{link}\n"
        f"{'-' * 68}\n\n{texto}\n"
    )


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--novos", type=int, default=10,
                    help="quantos primeiros contatos gerar (padrao: 10)")
    ap.add_argument("--marcar", action="store_true",
                    help="grava a data de hoje no Data_D0 dos novos")
    args = ap.parse_args()

    df = carregar()
    responderam, atrasados, hoje_lista, novos = classificar(df)
    novos_hoje = novos[:args.novos]

    hoje_str = dt.date.today().strftime("%d/%m/%Y")
    linhas = [
        "=" * 68,
        f"ROTEIRO DE PROSPECCAO - {hoje_str}",
        "Rodrigues Contabilidade",
        "=" * 68,
        "",
        f"  Responderam (responda PRIMEIRO) : {len(responderam)}",
        f"  Follow-ups ATRASADOS            : {len(atrasados)}",
        f"  Follow-ups de HOJE              : {len(hoje_lista)}",
        f"  Novos primeiros contatos        : {len(novos_hoje)}",
        f"  Fila restante na base           : {max(len(novos) - len(novos_hoje), 0)}",
        "",
    ]

    if responderam:
        linhas += ["", "#" * 68, "1. RESPONDERAM - NAO USE SCRIPT, RESPONDA VOCE",
                   "#" * 68]
        for l in responderam:
            linhas.append(bloco(
                l, ">> AGUARDANDO SUA RESPOSTA",
                f"Status atual: {l['Status']}\n"
                f"Abra a conversa e responda. Se pediu preco, NAO mande valor "
                f"antes do diagnostico.\nSe ja mandou os PGDAS, rode a analise "
                f"e entregue o relatorio de 1 pagina."))

    if atrasados:
        linhas += ["", "#" * 68, "2. ATRASADOS - ZERE ESTA FILA", "#" * 68]
        for l, etapa, func in atrasados:
            linhas.append(bloco(l, f">> {etapa} EM ATRASO", func(l)))

    if hoje_lista:
        linhas += ["", "#" * 68, "3. FOLLOW-UPS DE HOJE", "#" * 68]
        for l, etapa, func in hoje_lista:
            linhas.append(bloco(l, f">> {etapa}", func(l)))

    if novos_hoje:
        linhas += ["", "#" * 68, "4. NOVOS PRIMEIROS CONTATOS (D0)", "#" * 68]
        for l in novos_hoje:
            linhas.append(bloco(l, f">> D0 - {l['Cunha'] or 'Monofasico'}", d0(l)))
    else:
        linhas += ["", ">> FILA VAZIA. Rode o script 01 para repor a base."]

    linhas += ["", "=" * 68,
               "APOS ENVIAR: atualize Status e Data_D0 no CRM.",
               "Nunca mais de 3 mensagens em 7 dias para o mesmo contato.",
               "=" * 68]

    with open(SAIDA, "w", encoding="utf-8") as f:
        f.write("\n".join(linhas))

    print("\n".join(linhas[:12]))
    print(f"\nRoteiro completo gravado em: {SAIDA}")

    if args.marcar and novos_hoje:
        wb = load_workbook(CRM)
        ws = wb["Pipeline"]
        alvos = {l["Razao_Social"] for l in novos_hoje}
        n = 0
        for linha in range(3, ws.max_row + 1):
            if ws[f"C{linha}"].value in alvos and not ws[f"M{linha}"].value:
                ws[f"M{linha}"] = dt.date.today()
                ws[f"M{linha}"].number_format = "DD/MM/YYYY"
                ws[f"L{linha}"] = "D0 enviado"
                n += 1
        wb.save(CRM)
        print(f"CRM atualizado: {n} contatos marcados com D0 = {hoje_str}")


if __name__ == "__main__":
    main()
